"""
Resume Parser - DOCX Handler

Extracts text content from Microsoft Word (.docx) resume files.
"""

import io
import re
from typing import Optional, Tuple

try:
    from docx import Document
except ImportError:
    Document = None


def extract_text_from_docx(file_content: bytes) -> str:
    """
    Extract text content from a DOCX file.
    
    Preserves paragraph structure and handles various DOCX elements:
    - Paragraphs
    - Tables (converted to text)
    - Lists (bullet points preserved conceptually)
    
    Args:
        file_content: Raw bytes of the DOCX file
        
    Returns:
        Extracted text content as a string
        
    Raises:
        ValueError: If python-docx is not installed
        RuntimeError: If DOCX parsing fails
    """
    if Document is None:
        raise ValueError(
            "python-docx is required for DOCX parsing. "
            "Install it with: pip install python-docx"
        )
    
    try:
        # Create a file-like object from bytes
        docx_stream = io.BytesIO(file_content)
        doc = Document(docx_stream)
        
        all_text_parts = []
        
        # Extract text from paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                # Detect if this might be a heading (often has different styling)
                # DOCX paragraphs have style information we can use
                if para.style and para.style.name and 'Heading' in para.style.name:
                    # Add extra newline before headings for clarity
                    all_text_parts.append(f"\n{text}")
                else:
                    all_text_parts.append(text)
        
        # Extract text from tables
        table_text = _extract_table_text(doc)
        if table_text:
            all_text_parts.append("\n" + table_text)
        
        if not all_text_parts:
            return "[Warning: No text could be extracted from this DOCX file. " \
                   "The file may be empty or contain only images.]"
        
        return '\n'.join(all_text_parts)
    
    except Exception as e:
        raise RuntimeError(f"Failed to parse DOCX: {str(e)}") from e


def _extract_table_text(doc) -> str:
    """
    Extract text from tables in a DOCX document.
    
    Tables in resumes often contain skills or contact information.
    We extract them row by row.
    
    Args:
        doc: python-docx Document object
        
    Returns:
        Text content from all tables
    """
    table_parts = []
    
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                table_parts.append(' | '.join(row_text))
    
    return '\n'.join(table_parts)


def has_tables(file_content: bytes) -> bool:
    """
    Check if a DOCX contains tables.
    
    Args:
        file_content: Raw bytes of the DOCX file
        
    Returns:
        True if tables are present, False otherwise
    """
    if Document is None:
        return False
    
    try:
        docx_stream = io.BytesIO(file_content)
        doc = Document(docx_stream)
        return len(doc.tables) > 0
    except Exception:
        return False


def get_document_structure(file_content: bytes) -> dict:
    """
    Analyze the structure of a DOCX document.
    
    Useful for ATS analysis to understand document complexity.
    
    Args:
        file_content: Raw bytes of the DOCX file
        
    Returns:
        Dictionary with structure information
    """
    if Document is None:
        return {"error": "python-docx not installed"}
    
    try:
        docx_stream = io.BytesIO(file_content)
        doc = Document(docx_stream)
        
        return {
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "section_count": len(doc.sections),
            "has_headers": any(
                p.style and p.style.name and 'Heading' in p.style.name 
                for p in doc.paragraphs
            ),
        }
    except Exception as e:
        return {"error": str(e)}
